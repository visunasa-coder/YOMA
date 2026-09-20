from pathlib import Path
import json
import os
import zipfile

import pytest
from fastapi.testclient import TestClient

from yoma.app import create_app
from yoma.config import Settings
from yoma.db import connection_scope
from yoma.documents import DocumentError, extract_document, validate_document
from yoma.security import hash_password


def make_client(tmp_path: Path, *, max_document_size: int = 25 * 1024 * 1024, max_archive_members: int = 2000) -> TestClient:
    settings = Settings(
        database_path=tmp_path / "yoma.sqlite3",
        approved_roots=(),
        host="127.0.0.1",
        port=8765,
        session_ttl_seconds=3600,
        bootstrap_username="admin",
        bootstrap_password="correct horse battery staple",
        credential_vault_path=tmp_path / "credentials.vault",
        credential_vault_key="test-vault-key-" + ("x" * 32),
        max_document_size=max_document_size,
        max_archive_members=max_archive_members,
    )
    return TestClient(create_app(settings))


def login(client: TestClient, username: str = "admin", password: str = "correct horse battery staple") -> dict[str, str]:
    response = client.post("/auth/login", json={"username": username, "password": password})
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def approve_root(client: TestClient, headers: dict[str, str], root: Path) -> int:
    response = client.post("/workspace/roots", json={"path": str(root)}, headers=headers)
    assert response.status_code == 201
    return response.json()["id"]


def make_pdf(path: Path, text: str) -> None:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode())
        output.extend(obj)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    output.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]))
    output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(output)


def make_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("DOCX paragraph")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Table A"
    table.rows[0].cells[1].text = "Table B"
    document.save(path)


def make_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Summary"
    worksheet["A1"] = "XLSX value"
    workbook.save(path)


def test_supported_documents_extract_and_index_with_source_attribution(tmp_path: Path) -> None:
    root = tmp_path / "workspace"
    root.mkdir()
    (root / "notes.txt").write_text("TXT searchable phrase", encoding="utf-8")
    (root / "readme.md").write_text("# Markdown heading\nMarkdown searchable phrase", encoding="utf-8")
    make_pdf(root / "report.pdf", "PDF searchable phrase")
    make_docx(root / "letter.docx")
    make_xlsx(root / "sheet.xlsx")

    with make_client(tmp_path) as client:
        headers = login(client)
        root_id = approve_root(client, headers, root)
        document_ids = {}
        for filename in ("notes.txt", "readme.md", "report.pdf", "letter.docx", "sheet.xlsx"):
            response = client.post(
                "/documents/ingest",
                json={"workspace_root_id": root_id, "relative_path": filename},
                headers=headers,
            )
            assert response.status_code == 201, response.text
            document_ids[filename] = response.json()["document_id"]
            assert response.json()["workspace_root_id"] == root_id
            assert response.json()["relative_path"] == filename

        metadata = client.get(f"/documents/{document_ids['notes.txt']}", headers=headers)
        assert metadata.status_code == 200
        assert metadata.json()["content_type"] == "text/plain"
        assert metadata.json()["extraction_status"] == "extracted"

        text = client.get(f"/documents/{document_ids['notes.txt']}/text", headers=headers)
        assert text.status_code == 200
        assert "TXT searchable phrase" in text.json()["extracted_text"]

        results = client.get("/documents/search", params={"q": "searchable phrase"}, headers=headers)
        assert results.status_code == 200
        result_paths = {item["relative_path"] for item in results.json()}
        assert {"notes.txt", "readme.md", "report.pdf"}.issubset(result_paths)
        source = next(item for item in results.json() if item["relative_path"] == "notes.txt")
        assert source["document_id"] == document_ids["notes.txt"]
        assert source["workspace_root_id"] == root_id
        assert source["filename"] == "notes.txt"
        assert "TXT searchable phrase" in source["excerpt"]


def test_document_safety_rejects_unsupported_missing_directory_large_and_escape_paths(tmp_path: Path) -> None:
    root = tmp_path / "workspace"
    outside = tmp_path / "outside.txt"
    root.mkdir()
    outside.write_text("outside", encoding="utf-8")
    (root / "image.bin").write_bytes(b"binary")
    (root / "large.txt").write_bytes(b"12345")
    (root / "folder").mkdir()
    with make_client(tmp_path, max_document_size=4) as client:
        headers = login(client)
        root_id = approve_root(client, headers, root)
        assert client.post("/documents/ingest", json={"workspace_root_id": root_id, "relative_path": "image.bin"}, headers=headers).status_code == 415
        assert client.post("/documents/ingest", json={"workspace_root_id": root_id, "relative_path": "missing.txt"}, headers=headers).status_code == 400
        assert client.post("/documents/ingest", json={"workspace_root_id": root_id, "relative_path": "folder"}, headers=headers).status_code == 400
        assert client.post("/documents/ingest", json={"workspace_root_id": root_id, "relative_path": "large.txt"}, headers=headers).status_code == 400
        assert client.post("/documents/ingest", json={"workspace_root_id": root_id, "relative_path": "../outside.txt"}, headers=headers).status_code == 400
        absolute = os.fspath(outside.resolve())
        assert client.post("/documents/ingest", json={"workspace_root_id": root_id, "relative_path": absolute}, headers=headers).status_code == 400

        with connection_scope(tmp_path / "yoma.sqlite3") as connection:
            events = [row["event_type"] for row in connection.execute("SELECT event_type FROM audit_events").fetchall()]
        assert "document.unsupported" in events
        assert "document.unsafe_rejected" in events


def test_document_ownership_and_authentication_are_enforced(tmp_path: Path) -> None:
    root = tmp_path / "workspace"
    root.mkdir()
    (root / "private.txt").write_text("private phrase", encoding="utf-8")
    with make_client(tmp_path) as client:
        assert client.post("/documents/ingest", json={"workspace_root_id": 1, "relative_path": "private.txt"}).status_code == 401
        admin_headers = login(client)
        root_id = approve_root(client, admin_headers, root)
        document_id = client.post("/documents/ingest", json={"workspace_root_id": root_id, "relative_path": "private.txt"}, headers=admin_headers).json()["document_id"]
        with connection_scope(tmp_path / "yoma.sqlite3") as connection:
            connection.execute(
                "INSERT INTO users (username, password_hash, role, created_at) VALUES (?, ?, ?, datetime('now'))",
                ("second", hash_password("second password is long"), "user"),
            )
            connection.commit()
        second_headers = login(client, "second", "second password is long")
        assert client.get(f"/documents/{document_id}", headers=second_headers).status_code == 404
        assert client.get(f"/documents/{document_id}/text", headers=second_headers).status_code == 404
        assert client.get("/documents/search", params={"q": "private"}, headers=second_headers).json() == []


def test_malformed_documents_fail_safely_without_content_in_audit(tmp_path: Path) -> None:
    root = tmp_path / "workspace"
    root.mkdir()
    bad_pdf = root / "bad.pdf"
    bad_pdf.write_bytes(b"not a PDF secret phrase")
    bad_archive = root / "bad.docx"
    with zipfile.ZipFile(bad_archive, "w") as archive:
        for number in range(3):
            archive.writestr(f"member{number}.xml", "x")

    with make_client(tmp_path, max_archive_members=2) as client:
        headers = login(client)
        root_id = approve_root(client, headers, root)
        assert client.post("/documents/ingest", json={"workspace_root_id": root_id, "relative_path": "bad.pdf"}, headers=headers).status_code == 422
        assert client.post("/documents/ingest", json={"workspace_root_id": root_id, "relative_path": "bad.docx"}, headers=headers).status_code == 422
        with connection_scope(tmp_path / "yoma.sqlite3") as connection:
            statuses = connection.execute("SELECT extraction_status FROM documents ORDER BY document_id").fetchall()
            audit = " ".join(row["details_json"] for row in connection.execute("SELECT details_json FROM audit_events").fetchall())
        assert all(row["extraction_status"] in {"extraction_failed", "archive_member_limit"} for row in statuses)
        assert "secret phrase" not in audit


def test_document_validation_and_archive_limits_are_deterministic(tmp_path: Path) -> None:
    root = tmp_path / "workspace"
    root.mkdir()
    path = root / "note.txt"
    path.write_text("hello", encoding="utf-8")
    validated = validate_document(path, root, 1, 1, 100)
    assert validated.relative_path == "note.txt"
    assert "hello" in extract_document(path, 100, 100, 10).text
    with pytest.raises(DocumentError):
        validate_document(path, root, 1, 1, 2)

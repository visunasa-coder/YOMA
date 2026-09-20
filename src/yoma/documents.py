"""Bounded document validation, extraction, normalization, and identity."""

from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from pathlib import PurePosixPath, PureWindowsPath
import zipfile

from .workspace import WorkspacePathError, validate_child


SUPPORTED_EXTENSIONS = frozenset({".txt", ".md", ".pdf", ".docx", ".xlsx"})
CONTENT_TYPES = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


class DocumentError(ValueError):
    """Deterministic, safe document-processing failure."""

    def __init__(self, message: str, code: str):
        super().__init__(message)
        self.code = code


class DocumentLimitError(DocumentError):
    pass


@dataclass(frozen=True)
class ExtractionResult:
    text: str
    parser_name: str
    parser_version: str
    source_locations: tuple[dict[str, str | int], ...]


@dataclass(frozen=True)
class DocumentInput:
    document_id: str
    canonical_path: Path
    relative_path: str
    filename: str
    extension: str
    size: int
    modified_at: str
    content_type: str


def document_id_for(user_id: int, workspace_root_id: int, canonical_path: Path) -> str:
    value = f"{user_id}:{workspace_root_id}:{canonical_path}".encode("utf-8")
    return hashlib.sha256(value).hexdigest()


def validate_document(
    candidate: str | Path,
    root: str | Path,
    user_id: int,
    workspace_root_id: int,
    max_file_size: int,
) -> DocumentInput:
    try:
        canonical_path = validate_child(candidate, root)
    except WorkspacePathError as exc:
        raise DocumentError("document path is outside the approved root", "path_denied") from exc
    try:
        if not canonical_path.exists():
            raise DocumentError("document does not exist", "missing")
        if not canonical_path.is_file():
            raise DocumentError("document path is not a file", "not_a_file")
        metadata = canonical_path.stat()
    except OSError as exc:
        raise DocumentError("document is inaccessible", "inaccessible") from exc
    extension = canonical_path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentError("document extension is not supported", "unsupported_extension")
    if metadata.st_size > max_file_size:
        raise DocumentLimitError("document exceeds the configured size limit", "too_large")
    modified_at = datetime.fromtimestamp(metadata.st_mtime, timezone.utc).isoformat()
    return DocumentInput(
        document_id=document_id_for(user_id, workspace_root_id, canonical_path),
        canonical_path=canonical_path,
        relative_path=canonical_path.relative_to(Path(root).resolve(strict=False)).as_posix(),
        filename=canonical_path.name,
        extension=extension,
        size=metadata.st_size,
        modified_at=modified_at,
        content_type=CONTENT_TYPES[extension],
    )


def _bounded_text(text: str, maximum: int) -> str:
    if len(text.encode("utf-8")) > maximum:
        raise DocumentLimitError("extracted text exceeds the configured limit", "extracted_text_too_large")
    return text


def _safe_zip(path: Path, max_expansion: int, max_members: int) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > max_members:
                raise DocumentLimitError("archive contains too many members", "archive_member_limit")
            total_size = 0
            for member in members:
                normalized_name = member.filename.replace("\\", "/")
                parts = PurePosixPath(normalized_name).parts
                if (
                    PurePosixPath(normalized_name).is_absolute()
                    or PureWindowsPath(normalized_name).is_absolute()
                    or ".." in parts
                    or ".." in PureWindowsPath(normalized_name).parts
                    or ":" in PureWindowsPath(normalized_name).drive
                    or "\x00" in normalized_name
                ):
                    raise DocumentError("archive contains an unsafe member path", "unsafe_archive")
                total_size += member.file_size
                if total_size > max_expansion:
                    raise DocumentLimitError("archive expansion exceeds the configured limit", "archive_expansion_limit")
    except (zipfile.BadZipFile, OSError) as exc:
        raise DocumentError("document archive is malformed", "malformed_archive") from exc


def _extract_text(path: Path, maximum: int) -> ExtractionResult:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "utf-8"):
        try:
            return ExtractionResult(_bounded_text(data.decode(encoding), maximum), "text", "stdlib", ())
        except UnicodeDecodeError:
            continue
    return ExtractionResult(_bounded_text(data.decode("utf-8", errors="replace"), maximum), "text", "stdlib", ())


def _extract_pdf(path: Path, maximum: int) -> ExtractionResult:
    try:
        import pypdf

        reader = pypdf.PdfReader(str(path), strict=False)
        pieces: list[str] = []
        locations: list[dict[str, str | int]] = []
        for number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            pieces.append(f"Page {number}\n{text}")
            locations.append({"page": number})
            _bounded_text("\n\n".join(pieces), maximum)
        return ExtractionResult("\n\n".join(pieces), "pypdf", pypdf.__version__, tuple(locations))
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("PDF extraction failed", "extraction_failed") from exc


def _extract_docx(path: Path, maximum: int, max_expansion: int, max_members: int) -> ExtractionResult:
    _safe_zip(path, max_expansion, max_members)
    try:
        import docx

        document = docx.Document(str(path))
        pieces: list[str] = []
        locations: list[dict[str, str | int]] = []
        for number, paragraph in enumerate(document.paragraphs, start=1):
            if paragraph.text:
                pieces.append(paragraph.text)
                locations.append({"paragraph": number})
        for table_number, table in enumerate(document.tables, start=1):
            for row in table.rows:
                values = [cell.text for cell in row.cells]
                pieces.append(" | ".join(values))
                locations.append({"table": table_number})
        return ExtractionResult(_bounded_text("\n".join(pieces), maximum), "python-docx", docx.__version__, tuple(locations))
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("DOCX extraction failed", "extraction_failed") from exc


def _extract_xlsx(path: Path, maximum: int, max_expansion: int, max_members: int) -> ExtractionResult:
    _safe_zip(path, max_expansion, max_members)
    try:
        import openpyxl

        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True, keep_links=False)
        pieces: list[str] = []
        locations: list[dict[str, str | int]] = []
        for worksheet in workbook.worksheets:
            pieces.append(f"Sheet: {worksheet.title}")
            locations.append({"sheet": worksheet.title})
            for row in worksheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    pieces.append(" | ".join(values))
                    _bounded_text("\n".join(pieces), maximum)
        workbook.close()
        return ExtractionResult(_bounded_text("\n".join(pieces), maximum), "openpyxl", openpyxl.__version__, tuple(locations))
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("XLSX extraction failed", "extraction_failed") from exc


def extract_document(path: Path, max_text_size: int, max_archive_expansion: int, max_archive_members: int) -> ExtractionResult:
    try:
        if path.suffix.lower() in {".txt", ".md"}:
            return _extract_text(path, max_text_size)
        if path.suffix.lower() == ".pdf":
            return _extract_pdf(path, max_text_size)
        if path.suffix.lower() == ".docx":
            return _extract_docx(path, max_text_size, max_archive_expansion, max_archive_members)
        if path.suffix.lower() == ".xlsx":
            return _extract_xlsx(path, max_text_size, max_archive_expansion, max_archive_members)
    except DocumentError:
        raise
    except OSError as exc:
        raise DocumentError("document could not be read", "inaccessible") from exc
    raise DocumentError("document extension is not supported", "unsupported_extension")


def source_metadata(result: ExtractionResult) -> str:
    return json.dumps(list(result.source_locations), sort_keys=True)
